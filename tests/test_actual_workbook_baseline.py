from io import BytesIO
import hashlib
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from openpyxl.utils import get_column_letter

from core.excel_parser import parse_workbook
from core.normalizer import clean_text
from core.rich_text_parser import first_bold_author, read_rich_cells


WORKBOOK = Path(r"C:\Users\silic\OneDrive - Vardhaman College of Engineering\HoD\Publications_v2.xlsx")
EXPECTED_SHA256 = "641ECD8B9F40F136F105A97B0EDE06D9F1E9086B692E79CA4395F0CD0CF615C8"
SHEET_SPECS = {
    "Journal": {"title": 2, "authors": 3, "date": 6},
    "Conf": {"title": 3, "authors": 2, "date": 5},
    "Chapter": {"title": 3, "authors": 2, "date": 5},
    "Books": {"title": 3, "authors": 2, "date": 6},
}


@pytest.mark.skipif(not WORKBOOK.exists(), reason="Validated Publications_v2.xlsx is not available on this machine")
def test_current_workbook_source_baseline_and_reconciliation():
    try:
        workbook_bytes = WORKBOOK.read_bytes()
    except PermissionError:
        pytest.skip("Validated workbook is temporarily locked by Excel or OneDrive")

    assert hashlib.sha256(workbook_bytes).hexdigest().upper() == EXPECTED_SHA256

    normalized, metadata = parse_workbook(workbook_bytes)
    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=False)
    rich_cells = read_rich_cells(workbook_bytes)

    assert len(normalized) == 352
    assert normalized.groupby("Source Sheet").size().to_dict() == {
        "Books": 6,
        "Chapter": 22,
        "Conf": 186,
        "Journal": 138,
    }
    assert metadata["years"] == [2023, 2024, 2025, 2026]
    assert normalized["Publication ID"].is_unique

    expected_source_rows: set[tuple[str, int]] = set()
    for sheet_name, spec in SHEET_SPECS.items():
        sheet = workbook[sheet_name]
        for source_row in range(2, sheet.max_row + 1):
            source_title = clean_text(sheet.cell(source_row, spec["title"]).value)
            source_authors = clean_text(sheet.cell(source_row, spec["authors"]).value)
            if not source_title and not source_authors:
                continue
            expected_source_rows.add((sheet_name, source_row))
            record = normalized[
                normalized["Source Sheet"].eq(sheet_name) & normalized["Source Row"].eq(source_row)
            ]
            assert len(record) == 1
            row = record.iloc[0]
            assert row["Title"] == source_title
            assert row["Original Authors"] == source_authors
            source_date = sheet.cell(source_row, spec["date"]).value
            expected_year = source_date.year if hasattr(source_date, "year") else int(source_date)
            assert int(row["Calendar Year"]) == expected_year

    normalized_source_rows = set(
        zip(normalized["Source Sheet"], normalized["Source Row"].astype(int), strict=True)
    )
    assert normalized_source_rows == expected_source_rows

    multiple_bold_rows = 0
    for sheet_name in ("Journal", "Conf"):
        author_column = SHEET_SPECS[sheet_name]["authors"]
        for source_row in range(2, workbook[sheet_name].max_row + 1):
            coordinate = f"{get_column_letter(author_column)}{source_row}"
            rich_cell = rich_cells[sheet_name].get(coordinate)
            bold_runs = [run for run in rich_cell.runs if run.bold and run.text.strip()] if rich_cell else []
            if len(bold_runs) < 2:
                continue
            multiple_bold_rows += 1
            record = normalized[
                normalized["Source Sheet"].eq(sheet_name) & normalized["Source Row"].eq(source_row)
            ]
            assert len(record) == 1
            assert record.iloc[0]["Original Faculty Name"] == first_bold_author(rich_cell)

    assert multiple_bold_rows == 5
