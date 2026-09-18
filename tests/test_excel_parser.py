from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from core.author_mapper import UNMAPPED
from core.excel_parser import parse_workbook


def workbook_bytes(include_conf: bool = False) -> bytes:
    workbook = Workbook()
    journal = workbook.active
    journal.title = "Journal"
    journal.append(["S.No", "Title of the paper", "Authors", "Name of Journal", "Month and Year", "Indexing", "Quartile", "DOI"])
    bold = InlineFont(b=True)
    journal.append([
        1,
        "First paper",
        CellRichText(["External Author; ", TextBlock(bold, "Patil Mounica"), "; ", TextBlock(bold, "Natarajan Karuppiah")]),
        "Engineering Journal",
        "March 2024",
        "SCIE + Scopus",
        "Q1",
        "10.1/example",
    ])
    journal.append([2, "Ambiguous paper", "Author One; Author Two", "Engineering Journal", "bad date", "Other Index", "", ""])
    if include_conf:
        conf = workbook.create_sheet("Conf")
        conf.append(["Paper Title", "Authors", "Conference Name", "Year"])
        conf.append(["Conference paper", "S. Ravivarman", "Test Conference", 2025])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_end_to_end_rich_text_first_bold_wins():
    frame, metadata = parse_workbook(workbook_bytes())
    first = frame.iloc[0]
    assert first["Original Faculty Name"] == "Patil Mounica"
    assert first["First Bold Author"] == "Patil Mounica"
    assert first["Author Selection Method"] == "First Bold Author"
    assert first["Calendar Year"] == 2024
    assert first["SCI_SCIE_Flag"] == 1 and first["Scopus_Flag"] == 1
    assert first["Q1_Flag"] == 1


def test_ambiguous_no_bold_row_is_unmapped():
    frame, _ = parse_workbook(workbook_bytes())
    second = frame.iloc[1]
    assert second["Original Faculty Name"] == UNMAPPED
    assert second["Date Status"] == "Invalid"


def test_missing_sheet_handling_warns_and_continues():
    frame, metadata = parse_workbook(workbook_bytes(include_conf=False))
    assert len(frame) == 2
    assert metadata["detected_sheets"] == ["Journal"]
    assert any("Missing optional publication sheets" in warning for warning in metadata["warnings"])


def test_no_bold_single_author_in_conference():
    frame, _ = parse_workbook(workbook_bytes(include_conf=True))
    record = frame[frame["Publication Type"] == "Conference"].iloc[0]
    assert record["Original Faculty Name"] == "S. Ravivarman"
    assert record["Author Selection Method"] == "No Bold / Single Author"


def test_actual_workbook_header_variants_for_chapter_and_book():
    workbook = Workbook()
    chapter = workbook.active
    chapter.title = "Chapter"
    chapter.append(["S.No", "Authors", "Title of the chapter", "Name of the book", "Month and Year", "ISBN", "DoI", "Indexing"])
    chapter.append([1, "Patil Mounica", "Chapter title", "Collected Work", 2025, "123", "10.1/chapter", "Scopus"])
    books = workbook.create_sheet("Books")
    books.append(["S.No", "Authors/Editors", "Name of the book", "Volume", "Issue", "Month and Year", "ISBN", "Name of Publishing Agency"])
    books.append([1, "Patil Mounica", "Book title", "-", "-", 2024, "456", "Test Publisher"])
    buffer = BytesIO()
    workbook.save(buffer)
    frame, _ = parse_workbook(buffer.getvalue())
    chapter_row = frame[frame["Publication Type"] == "Book Chapter"].iloc[0]
    book_row = frame[frame["Publication Type"] == "Book"].iloc[0]
    assert chapter_row["Title"] == "Chapter title" and chapter_row["Journal / Outlet Name"] == "Collected Work"
    assert book_row["Title"] == "Book title" and book_row["Publisher"] == "Test Publisher"
