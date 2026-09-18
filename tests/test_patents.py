from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
import pandas as pd

from core.author_mapper import apply_author_mappings
from core.excel_parser import parse_workbook
from core.patent_metrics import filter_patents, patent_kpis
from core.patent_parser import (
    add_patent_duplicate_flags,
    normalize_patent_attributions,
    normalize_patent_number,
    parse_patents,
)
from core.patent_state import (
    apply_patent_attribution_resolutions,
    apply_patent_duplicate_decisions,
    load_patent_attribution_resolutions,
    save_patent_attribution_resolutions,
)


BOLD = InlineFont(b=True)


def _workbook_bytes(include_patents: bool = True) -> bytes:
    workbook = Workbook()
    journal = workbook.active
    journal.title = "Journal"
    journal.append(["S.No", "Title", "Authors", "Month and Year"])
    journal.append(
        [
            1,
            "Publication",
            CellRichText([TextBlock(BOLD, "Faculty A"), "; ", TextBlock(BOLD, "Faculty B")]),
            2025,
        ]
    )
    if include_patents:
        patents = workbook.create_sheet("Patents")
        patents.append(
            ["Sl.No.", "Name of Faculty", "Title of the Patent", "Patent No", "Type of Patent", "Publication date", "Granted date"]
        )
        patents.append(
            [1, CellRichText(["External\n", TextBlock(BOLD, "Faculty A")]), "Single faculty", "PAT-001", "Utility", pd.Timestamp("2024-03-10"), None]
        )
        patents.append(
            [
                2,
                CellRichText(
                    [TextBlock(BOLD, "Faculty A"), "\n", TextBlock(BOLD, "Faculty B"), "\n", TextBlock(BOLD, "Faculty C")]
                ),
                "Joint patent",
                "PAT-002",
                "Design",
                pd.Timestamp("2025-04-11"),
                pd.Timestamp("2026-01-12"),
            ]
        )
        patents.append([3, "Faculty A", "Fallback patent", "PAT-003", "Utility", pd.Timestamp("2025-05-01"), None])
        patents.append([4, "Faculty A\nFaculty B", "Review patent", "PAT-004", "Utility", pd.Timestamp("2025-06-01"), None])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_patent_sheet_absent_keeps_dashboard_parser_functional():
    data = _workbook_bytes(include_patents=False)
    publications, _ = parse_workbook(data)
    patents, attribution, metadata = parse_patents(data)
    assert len(publications) == 1
    assert patents.empty and attribution.empty
    assert metadata["detected"] is False


def test_patent_header_variants_are_detected():
    master, _, metadata = parse_patents(_workbook_bytes())
    assert len(master) == 4
    assert metadata["headers"]["faculty"] == "Name of Faculty"
    assert metadata["headers"]["number"] == "Patent No"


def test_one_bold_patent_faculty_receives_one_credit():
    master, attribution, _ = parse_patents(_workbook_bytes())
    patent_id = master.loc[master["Title_of_Patent"].eq("Single faculty"), "Patent_ID"].iloc[0]
    rows = attribution[attribution["Patent_ID"].eq(patent_id)]
    assert rows["Original_Extracted_Name"].tolist() == ["Faculty A"]
    assert rows["Faculty_Patent_Credit"].sum() == 1


def test_multiple_bold_patent_faculty_all_receive_credit_without_institutional_double_count():
    master, attribution, _ = parse_patents(_workbook_bytes())
    joint = master[master["Title_of_Patent"].eq("Joint patent")]
    rows = attribution[attribution["Patent_ID"].isin(joint["Patent_ID"])]
    assert rows["Original_Extracted_Name"].tolist() == ["Faculty A", "Faculty B", "Faculty C"]
    assert rows["Faculty_Patent_Credit"].sum() == 3
    assert joint["Institutional_Patent_Count"].sum() == 1


def test_publication_multiple_bold_logic_remains_first_bold_only():
    publications, _ = parse_workbook(_workbook_bytes())
    assert len(publications) == 1
    assert publications.iloc[0]["Original Faculty Name"] == "Faculty A"


def test_patent_no_bold_single_faculty_fallback_and_multi_person_review():
    master, attribution, _ = parse_patents(_workbook_bytes())
    fallback = master[master["Title_of_Patent"].eq("Fallback patent")].iloc[0]
    review = master[master["Title_of_Patent"].eq("Review patent")].iloc[0]
    assert fallback["Patent_Attribution_Method"] == "No Bold / Single Faculty"
    assert fallback["Patent_ID"] in set(attribution["Patent_ID"])
    assert review["Attribution_Status"] == "REVIEW REQUIRED"
    assert review["Patent_ID"] not in set(attribution["Patent_ID"])


def test_patent_faculty_uses_publication_canonical_normalization():
    publication = pd.DataFrame(
        {
            "Publication ID": ["P1", "P2"],
            "Publication Type": ["Journal", "Journal"],
            "Original Faculty Name": ["N. Karuppiah", "Karuppiah Natarajan"],
        }
    )
    publication = apply_author_mappings(publication, {})
    attribution = pd.DataFrame(
        {
            "Patent_ID": ["PAT1"],
            "Faculty_Name": ["Natarajan Karuppiah"],
            "Original_Extracted_Name": ["Natarajan Karuppiah"],
            "Canonical_Faculty_Name": ["Natarajan Karuppiah"],
            "Mapping_Status": ["Pending"],
            "Attribution_Method": ["All Bold Faculty"],
            "Bold_Order": [1],
            "Source_Row": [2],
            "Faculty_Patent_Credit": [1],
        }
    )
    result = normalize_patent_attributions(attribution, publication, {})
    assert result.iloc[0]["Canonical_Faculty_Name"] == "Karuppiah Natarajan"
    assert result.iloc[0]["Mapping_Status"] != "Review Required"


def test_two_patents_keep_distinct_institutional_and_faculty_credit_totals():
    master = pd.DataFrame(
        {
            "Patent_ID": ["PAT001", "PAT002"],
            "Publication_Date": [pd.Timestamp("2025-01-01")] * 2,
            "Granted_Date": [pd.NaT, pd.NaT],
            "Attribution_Status": ["Attributed", "Attributed"],
        }
    )
    attribution = pd.DataFrame(
        {"Patent_ID": ["PAT001", "PAT001", "PAT002"], "Faculty_Name": ["A", "B", "A"]}
    )
    values = patent_kpis(master, attribution)
    assert values["Total Patents"] == 2
    assert attribution.groupby("Faculty_Name")["Patent_ID"].nunique().to_dict() == {"A": 2, "B": 1}
    assert len(attribution) == 3


def test_publication_and_granted_dates_and_statuses_remain_separate():
    master, _, _ = parse_patents(_workbook_bytes())
    published = master[master["Title_of_Patent"].eq("Single faculty")].iloc[0]
    granted = master[master["Title_of_Patent"].eq("Joint patent")].iloc[0]
    assert published["Publication_Year"] == 2024
    assert published["Granted_Date_Status"] == "Missing"
    assert published["Patent_Status"] == "Published"
    assert granted["Publication_Year"] == 2025
    assert granted["Granted_Year"] == 2026
    assert granted["Patent_Status"] == "Granted"


def test_patent_number_normalization_and_duplicate_detection():
    assert normalize_patent_number(" pat-001 ") == normalize_patent_number("PAT 001")
    master = pd.DataFrame(
        {
            "Patent_Number_Normalized": [normalize_patent_number("pat-001"), normalize_patent_number("PAT 001")],
            "Title_of_Patent": ["Title One", "Different Title"],
        }
    )
    result = add_patent_duplicate_flags(master)
    assert set(result["Patent_Duplicate_Classification"]) == {"Strong Duplicate"}


def test_same_title_with_different_numbers_requires_review():
    master = pd.DataFrame(
        {
            "Patent_Number_Normalized": ["PAT001", "PAT002"],
            "Title_of_Patent": ["Shared Patent Title", "shared patent-title"],
        }
    )
    result = add_patent_duplicate_flags(master)
    assert set(result["Patent_Duplicate_Classification"]) == {"Review Required"}


def test_manual_multi_faculty_resolution_persists_and_creates_multiple_credits():
    master, attribution, _ = parse_patents(_workbook_bytes())
    patent_id = master.loc[master["Title_of_Patent"].eq("Review patent"), "Patent_ID"].iloc[0]
    with TemporaryDirectory() as directory:
        path = Path(directory) / "resolutions.json"
        save_patent_attribution_resolutions({patent_id: ["Faculty A", "Faculty B"]}, path)
        saved = load_patent_attribution_resolutions(path)
    resolved_master, resolved_attribution = apply_patent_attribution_resolutions(master, attribution, saved)
    credits = resolved_attribution[resolved_attribution["Patent_ID"].eq(patent_id)]
    assert set(credits["Faculty_Name"]) == {"Faculty A", "Faculty B"}
    assert credits["Faculty_Patent_Credit"].sum() == 2
    assert resolved_master.loc[resolved_master["Patent_ID"].eq(patent_id), "Attribution_Status"].iloc[0] == "Manually Resolved"


def test_explicit_patent_duplicate_exclusion_and_filters():
    master, attribution, _ = parse_patents(_workbook_bytes())
    excluded_id = master.iloc[0]["Patent_ID"]
    master.loc[master["Patent_ID"].eq(excluded_id), "Possible_Patent_Duplicate"] = "Yes"
    master = apply_patent_duplicate_decisions(master, {excluded_id: "Exclude"})
    filtered_master, filtered_attribution = filter_patents(master, attribution, include_duplicates=False)
    assert excluded_id not in set(filtered_master["Patent_ID"])
    assert excluded_id not in set(filtered_attribution["Patent_ID"])
    year_master, _ = filter_patents(master, attribution, publication_years=[2024])
    assert set(year_master["Publication_Year"]) == {2024}


def test_patent_grant_faculty_type_and_status_filters_apply_together():
    master, attribution, _ = parse_patents(_workbook_bytes())
    filtered_master, filtered_attribution = filter_patents(
        master,
        attribution,
        granted_years=[2026],
        faculty=["Faculty A"],
        patent_types=["Design"],
        statuses=["Granted"],
    )
    assert filtered_master["Title_of_Patent"].tolist() == ["Joint patent"]
    assert set(filtered_attribution["Faculty_Name"]) == {"Faculty A"}


def test_adding_patents_does_not_change_publication_row_count():
    without_patents, _ = parse_workbook(_workbook_bytes(include_patents=False))
    with_patents, _ = parse_workbook(_workbook_bytes(include_patents=True))
    assert len(without_patents) == len(with_patents) == 1


def test_patent_excel_export_contains_both_logical_datasets():
    from core.patent_metrics import patent_excel_bytes

    master, attribution, _ = parse_patents(_workbook_bytes())
    workbook = openpyxl.load_workbook(BytesIO(patent_excel_bytes(master, attribution)), read_only=True)
    assert workbook.sheetnames == ["Patent Master", "Patent Faculty Attribution"]
