from __future__ import annotations

from collections import Counter
from io import BytesIO
import hashlib
import re
from typing import Any

import openpyxl
import pandas as pd

from .author_mapper import UNMAPPED, _could_match, apply_author_mappings, author_key
from .normalizer import clean_text, normalized_key, parse_year
from .rich_text_parser import RichCell, looks_like_single_author, read_rich_cells


PATENT_SHEET_NAMES = {"patent", "patents"}
PATENT_ALIASES = {
    "serial": ["slno", "sno", "serialno", "serialnumber"],
    "faculty": ["nameofthefaculty", "nameoffaculty", "facultyname", "facultyapplicants", "applicants"],
    "title": ["titleofthepatent", "patenttitle", "title"],
    "number": ["patentnumber", "patentno", "applicationnumber", "applicationno"],
    "type": ["typeofpatent", "patenttype", "type"],
    "publication_date": ["publicationdate", "patentpublicationdate", "publisheddate"],
    "granted_date": ["granteddate", "grantdate", "dateofgrant"],
}

PATENT_MASTER_COLUMNS = [
    "Patent_ID",
    "Source_Sheet",
    "Source_Row",
    "Source_SNo",
    "Title_of_Patent",
    "Patent_Number_Raw",
    "Patent_Number_Normalized",
    "Patent_Type",
    "Publication_Date",
    "Publication_Year",
    "Publication_Date_Status",
    "Granted_Date",
    "Granted_Year",
    "Granted_Date_Status",
    "Patent_Status",
    "Institutional_Patent_Count",
    "Possible_Patent_Duplicate",
    "Patent_Duplicate_Classification",
    "Patent_Duplicate_Reason",
    "Patent_Duplicate_Key",
    "Patent_Duplicate_Review_Decision",
    "Original_Faculty_Cell",
    "Bold_Names_Detected",
    "Patent_Attribution_Method",
    "Attribution_Status",
]

PATENT_ATTRIBUTION_COLUMNS = [
    "Patent_ID",
    "Faculty_Name",
    "Original_Extracted_Name",
    "Canonical_Faculty_Name",
    "Mapping_Status",
    "Attribution_Method",
    "Bold_Order",
    "Source_Row",
    "Faculty_Patent_Credit",
]


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(value).lower())


def _field_for_header(value: object) -> str | None:
    key = _header_key(value)
    if not key:
        return None
    for field, aliases in PATENT_ALIASES.items():
        if key in aliases:
            return field
    if "faculty" in key and ("name" in key or "applicant" in key):
        return "faculty"
    if "patent" in key and "title" in key:
        return "title"
    if "patent" in key and ("number" in key or key.endswith("no")):
        return "number"
    if "patent" in key and "type" in key:
        return "type"
    if "publication" in key and "date" in key:
        return "publication_date"
    if ("grant" in key or "granted" in key) and "date" in key:
        return "granted_date"
    return None


def _detect_header_row(sheet: openpyxl.worksheet.worksheet.Worksheet) -> tuple[int, dict[str, int]]:
    best_row, best_fields, best_score = 1, {}, -1
    for row in range(1, min(sheet.max_row, 20) + 1):
        fields: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            field = _field_for_header(sheet.cell(row, column).value)
            if field and field not in fields:
                fields[field] = column
        score = len(fields) + (2 if "title" in fields else 0) + (2 if "faculty" in fields else 0)
        if score > best_score:
            best_row, best_fields, best_score = row, fields, score
    if "title" not in best_fields or "faculty" not in best_fields:
        raise ValueError("The Patents sheet must contain recognizable patent-title and faculty-name headers.")
    return best_row, best_fields


def _value(sheet: openpyxl.worksheet.worksheet.Worksheet, row: int, fields: dict[str, int], field: str) -> Any:
    column = fields.get(field)
    return sheet.cell(row, column).value if column else None


def _clean_faculty_name(value: str) -> str:
    value = re.sub(r"^\s*\d+\s*[.)-]\s*", "", value)
    return re.sub(r"^[\s,;|:&-]+|[\s,;|:&-]+$", "", re.sub(r"\s+", " ", value)).strip()


def _faculty_cell_text(value: object) -> str:
    if value is None:
        return ""
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in str(value).replace("\u00a0", " ").splitlines()]
    return "\n".join(line for line in lines if line)


def all_bold_faculty(cell: RichCell | None) -> list[str]:
    """Return every non-empty bold faculty segment in reading order."""
    if cell is None:
        return []
    names: list[str] = []
    for run in cell.runs:
        if not run.bold or not run.text.strip():
            continue
        for segment in re.split(r"\r?\n|;|\||\s+and\s+|\s*&\s*", run.text, flags=re.IGNORECASE):
            name = _clean_faculty_name(segment)
            if name and name not in names:
                names.append(name)
    return names


def normalize_patent_number(value: object) -> str:
    """Normalize harmless case, whitespace and punctuation differences for matching."""
    return re.sub(r"[^A-Z0-9]", "", clean_text(value).upper())


def _patent_id(sheet: str, row: int, title: object, number: object) -> str:
    seed = f"{sheet}|{row}|{clean_text(title)}|{clean_text(number)}"
    return "PAT" + hashlib.sha1(seed.encode("utf-8")).hexdigest()[:12].upper()


def add_patent_duplicate_flags(master: pd.DataFrame) -> pd.DataFrame:
    result = master.copy()
    result["Possible_Patent_Duplicate"] = "No"
    result["Patent_Duplicate_Classification"] = "Not Duplicate"
    result["Patent_Duplicate_Reason"] = ""
    result["Patent_Duplicate_Key"] = ""

    for number, group in result[result["Patent_Number_Normalized"].ne("")].groupby("Patent_Number_Normalized"):
        if len(group) < 2:
            continue
        result.loc[
            group.index,
            [
                "Possible_Patent_Duplicate",
                "Patent_Duplicate_Classification",
                "Patent_Duplicate_Reason",
                "Patent_Duplicate_Key",
            ],
        ] = ["Yes", "Strong Duplicate", "Same normalized patent number", f"NUMBER:{number}"]

    title_keys = result["Title_of_Patent"].map(normalized_key)
    for title, group in result[title_keys.ne("")].groupby(title_keys[title_keys.ne("")]):
        if len(group) < 2:
            continue
        unclassified = group.index[result.loc[group.index, "Patent_Duplicate_Classification"].eq("Not Duplicate")]
        if len(unclassified) == 0:
            continue
        result.loc[
            unclassified,
            [
                "Possible_Patent_Duplicate",
                "Patent_Duplicate_Classification",
                "Patent_Duplicate_Reason",
                "Patent_Duplicate_Key",
            ],
        ] = ["Yes", "Review Required", "Same normalized patent title", f"TITLE:{title}"]
    return result


def parse_patents(file_bytes: bytes) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    """Parse the optional Patents sheet into master and faculty-attribution datasets."""
    workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    sheet_name = next((name for name in workbook.sheetnames if name.strip().lower() in PATENT_SHEET_NAMES), None)
    if sheet_name is None:
        return (
            pd.DataFrame(columns=PATENT_MASTER_COLUMNS),
            pd.DataFrame(columns=PATENT_ATTRIBUTION_COLUMNS),
            {"detected": False, "sheet": None, "headers": {}, "warnings": [], "processed_rows": 0},
        )

    sheet = workbook[sheet_name]
    header_row, fields = _detect_header_row(sheet)
    rich_cells = read_rich_cells(file_bytes).get(sheet_name, {})
    records: list[dict[str, object]] = []
    attributions: list[dict[str, object]] = []

    for row in range(header_row + 1, sheet.max_row + 1):
        title = clean_text(_value(sheet, row, fields, "title"))
        faculty_cell = _faculty_cell_text(_value(sheet, row, fields, "faculty"))
        number = clean_text(_value(sheet, row, fields, "number"))
        if not title and not faculty_cell and not number:
            continue
        patent_id = _patent_id(sheet_name, row, title, number)
        rich_cell = rich_cells.get(f"{openpyxl.utils.get_column_letter(fields['faculty'])}{row}")
        bold_names = all_bold_faculty(rich_cell)
        if bold_names:
            extracted_names = bold_names
            attribution_method = "All Bold Faculty"
            attribution_status = "Attributed"
        elif looks_like_single_author(faculty_cell):
            extracted_names = [faculty_cell]
            attribution_method = "No Bold / Single Faculty"
            attribution_status = "Attributed"
        else:
            extracted_names = []
            attribution_method = "No Bold / Multiple People"
            attribution_status = "REVIEW REQUIRED"

        publication_year, publication_date, publication_date_status = parse_year(
            _value(sheet, row, fields, "publication_date")
        )
        granted_year, granted_date, granted_date_status = parse_year(_value(sheet, row, fields, "granted_date"))
        if granted_date_status == "Valid":
            patent_status = "Granted"
        elif publication_date_status == "Valid":
            patent_status = "Published"
        else:
            patent_status = "Status Not Available"

        records.append(
            {
                "Patent_ID": patent_id,
                "Source_Sheet": sheet_name,
                "Source_Row": row,
                "Source_SNo": _value(sheet, row, fields, "serial"),
                "Title_of_Patent": title,
                "Patent_Number_Raw": number,
                "Patent_Number_Normalized": normalize_patent_number(number),
                "Patent_Type": clean_text(_value(sheet, row, fields, "type")),
                "Publication_Date": publication_date,
                "Publication_Year": publication_year,
                "Publication_Date_Status": publication_date_status,
                "Granted_Date": granted_date,
                "Granted_Year": granted_year,
                "Granted_Date_Status": granted_date_status,
                "Patent_Status": patent_status,
                "Institutional_Patent_Count": 1,
                "Possible_Patent_Duplicate": "No",
                "Patent_Duplicate_Classification": "Not Duplicate",
                "Patent_Duplicate_Reason": "",
                "Patent_Duplicate_Key": "",
                "Patent_Duplicate_Review_Decision": "Not Applicable",
                "Original_Faculty_Cell": faculty_cell,
                "Bold_Names_Detected": " | ".join(bold_names),
                "Patent_Attribution_Method": attribution_method,
                "Attribution_Status": attribution_status,
            }
        )
        for order, name in enumerate(extracted_names, start=1):
            attributions.append(
                {
                    "Patent_ID": patent_id,
                    "Faculty_Name": name,
                    "Original_Extracted_Name": name,
                    "Canonical_Faculty_Name": name,
                    "Mapping_Status": "Pending",
                    "Attribution_Method": attribution_method,
                    "Bold_Order": order if bold_names else None,
                    "Source_Row": row,
                    "Faculty_Patent_Credit": 1,
                }
            )

    master = add_patent_duplicate_flags(pd.DataFrame(records, columns=PATENT_MASTER_COLUMNS))
    attribution = pd.DataFrame(attributions, columns=PATENT_ATTRIBUTION_COLUMNS)
    metadata = {
        "detected": True,
        "sheet": sheet_name,
        "headers": {field: sheet.cell(header_row, column).value for field, column in fields.items()},
        "warnings": [],
        "processed_rows": len(master),
        "publication_years": sorted(int(value) for value in master["Publication_Year"].dropna().unique()),
        "granted_years": sorted(int(value) for value in master["Granted_Year"].dropna().unique()),
    }
    return master, attribution, metadata


def normalize_patent_attributions(
    attribution: pd.DataFrame,
    publication_frame: pd.DataFrame,
    saved_mapping: dict[str, str] | None = None,
) -> pd.DataFrame:
    """Reuse publication canonical identities without altering publication attribution."""
    if attribution.empty:
        return attribution.copy()
    saved_mapping = saved_mapping or {}
    publication_seed = publication_frame.loc[
        publication_frame["Original Faculty Name"].ne(UNMAPPED),
        ["Publication ID", "Publication Type", "Original Faculty Name"],
    ].copy()
    patent_seed = pd.DataFrame(
        {
            "Publication ID": [f"__PATENT_{position}" for position in range(len(attribution))],
            "Publication Type": "Patent",
            "Original Faculty Name": attribution["Original_Extracted_Name"].values,
        }
    )
    combined = pd.concat([publication_seed, patent_seed], ignore_index=True)
    normalized = apply_author_mappings(combined, saved_mapping).tail(len(patent_seed))
    result = attribution.copy()
    result["Canonical_Faculty_Name"] = normalized["Faculty Name"].values
    result["Faculty_Name"] = result["Canonical_Faculty_Name"]
    result["Mapping_Status"] = normalized["Mapping Status"].values

    publication_aliases = publication_frame.loc[
        publication_frame["Original Faculty Name"].ne(UNMAPPED),
        ["Original Faculty Name", "Faculty Name"],
    ].drop_duplicates()
    for index, row in result.iterrows():
        original = clean_text(row["Original_Extracted_Name"])
        candidates = {
            clean_text(candidate["Faculty Name"])
            for _, candidate in publication_aliases.iterrows()
            if clean_text(candidate["Faculty Name"]) != UNMAPPED
            and (
                author_key(original) == author_key(candidate["Original Faculty Name"])
                or _could_match(original, clean_text(candidate["Original Faculty Name"]))
            )
        }
        if len(candidates) == 1:
            canonical = candidates.pop()
            result.at[index, "Canonical_Faculty_Name"] = canonical
            result.at[index, "Faculty_Name"] = canonical
            if canonical != original:
                result.at[index, "Mapping_Status"] = "Auto-normalized"

    canonical_roster = {
        clean_text(value)
        for value in publication_frame["Faculty Name"].dropna().unique()
        if clean_text(value) and clean_text(value) != UNMAPPED
    }
    approved_targets = {clean_text(value) for value in saved_mapping.values() if clean_text(value)}
    recognized = result["Canonical_Faculty_Name"].isin(canonical_roster | approved_targets)
    result.loc[~recognized, "Mapping_Status"] = "Review Required"
    return result[PATENT_ATTRIBUTION_COLUMNS]
