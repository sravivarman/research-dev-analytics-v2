from __future__ import annotations

from io import BytesIO
import re
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from .author_mapper import UNMAPPED
from .normalizer import clean_text, indexing_flags, normalize_quartile, parse_year, publication_id
from .rich_text_parser import RichCell, first_bold_author, looks_like_single_author, read_rich_cells


SHEET_TYPES = {
    "journal": "Journal",
    "conf": "Conference",
    "conference": "Conference",
    "chapter": "Book Chapter",
    "book chapter": "Book Chapter",
    "books": "Book",
    "book": "Book",
}

ALIASES = {
    "serial": ["sno", "slno", "serialno", "serialnumber"],
    "title": ["titleofthepaper", "papertitle", "title", "booktitle", "titleofthebook", "titleofthechapter", "chaptertitle"],
    "authors": ["authors", "author", "nameofauthors", "authorname", "facultyname", "faculty"],
    "outlet": ["nameofjournal", "nameofthejournal", "journalname", "conference name", "nameofconference", "bookpublicationname", "nameofthebook", "bookname", "publicationname", "outlet"],
    "doi": ["doi", "doilink", "doiurl"],
    "publisher": ["publisher", "publishername", "nameofpublisher", "nameofpublishingagency", "publishingagency"],
    "volume_issue": ["volumeissue", "volumeissueno", "volume", "volissue"],
    "issue": ["issue", "issueno"],
    "date": ["monthandyear", "monthyear", "publicationdate", "date", "yearofpublication", "publicationyear", "year"],
    "indexing": ["indexing", "indexed in", "indexedin", "indexingdetails", "database"],
    "impact_factor": ["impactfactor", "impactfactorif", "if"],
    "quartile": ["quartile", "journalquartile", "qranking"],
    "isbn": ["isbnissn", "isbn", "issn", "isbnno", "issnno"],
}

EXPECTED_FIELDS = set(ALIASES)


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(value).lower())


def _field_for_header(value: object) -> str | None:
    key = _header_key(value)
    if not key:
        return None
    for field, aliases in ALIASES.items():
        normalized = {_header_key(alias) for alias in aliases}
        if key in normalized:
            return field
    # Conservative partial matches for verbose real-world headers.
    if "author" in key or "facultyname" in key:
        return "authors"
    if "quartile" in key:
        return "quartile"
    if "impactfactor" in key:
        return "impact_factor"
    if "index" in key and len(key) < 40:
        return "indexing"
    if "month" in key and "year" in key:
        return "date"
    if key.startswith("doi"):
        return "doi"
    if "publisher" in key:
        return "publisher"
    if "isbn" in key or "issn" in key:
        return "isbn"
    return None


def _detect_header_row(sheet: openpyxl.worksheet.worksheet.Worksheet) -> tuple[int, dict[str, int]]:
    best_row, best_map, best_score = 1, {}, -1
    for row in range(1, min(sheet.max_row, 20) + 1):
        mapping: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            field = _field_for_header(sheet.cell(row, column).value)
            if field and field not in mapping:
                mapping[field] = column
        score = len(mapping) + (2 if "title" in mapping else 0) + (2 if "authors" in mapping else 0)
        if score > best_score:
            best_row, best_map, best_score = row, mapping, score
    if "title" not in best_map and "authors" not in best_map:
        raise ValueError(f"Could not identify a publication header row in sheet '{sheet.title}'.")
    return best_row, best_map


def _value(sheet: openpyxl.worksheet.worksheet.Worksheet, row: int, mapping: dict[str, int], field: str) -> Any:
    column = mapping.get(field)
    return sheet.cell(row, column).value if column else None


def _rich_at(rich: dict[str, RichCell], row: int, column: int | None) -> RichCell | None:
    return rich.get(f"{get_column_letter(column)}{row}") if column else None


def _select_author(publication_type: str, author_text: str, rich_cell: RichCell | None) -> tuple[str, str, str]:
    bold = first_bold_author(rich_cell)
    if bold:
        return bold, bold, "First Bold Author"
    if looks_like_single_author(author_text):
        return author_text, "", "No Bold / Single Author"
    return UNMAPPED, "", "Unmapped / Review Required"


def _safe_float(value: object) -> float | None:
    if value is None or clean_text(value) == "":
        return None
    match = re.search(r"[-+]?\d+(?:\.\d+)?", clean_text(value).replace(",", ""))
    try:
        return float(match.group()) if match else None
    except ValueError:
        return None


def _resolve_sheet_type(name: str) -> str | None:
    normalized = re.sub(r"\s+", " ", name.strip().lower())
    if normalized in SHEET_TYPES:
        return SHEET_TYPES[normalized]
    if normalized.startswith("journal"):
        return "Journal"
    if normalized.startswith("conf"):
        return "Conference"
    if "chapter" in normalized:
        return "Book Chapter"
    if normalized.startswith("book"):
        return "Book"
    return None


MASTER_COLUMNS = [
    "Publication ID", "Faculty Name", "Original Faculty Name", "Mapping Status", "Calendar Year",
    "Publication Date", "Date Status", "Publication Type", "Title", "Journal / Outlet Name", "DOI",
    "Publisher", "Volume / Issue", "ISBN/ISSN", "Indexing Raw", "SCI_SCIE_Flag", "ESCI_Flag",
    "Scopus_Flag", "Non_Scopus_Flag", "Indexing Recognized", "Quartile", "Q1_Flag", "Q2_Flag",
    "Q3_Flag", "Q4_Flag", "Impact Factor", "Original Authors", "First Bold Author",
    "Author Selection Method", "Source Sheet", "Source Row",
]


def parse_workbook(file_bytes: bytes) -> tuple[pd.DataFrame, dict[str, object]]:
    """Parse publication sheets into one row-based normalized master table."""
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError(f"Unable to open workbook: {exc}") from exc
    rich_by_sheet = read_rich_cells(file_bytes)
    detected: list[str] = []
    detected_types: set[str] = set()
    warnings: list[str] = []
    records: list[dict[str, object]] = []

    for sheet_name in workbook.sheetnames:
        publication_type = _resolve_sheet_type(sheet_name)
        if publication_type is None:
            continue
        detected.append(sheet_name)
        detected_types.add(publication_type)
        sheet = workbook[sheet_name]
        try:
            header_row, fields = _detect_header_row(sheet)
        except ValueError as exc:
            warnings.append(str(exc))
            continue
        if publication_type == "Book" and "title" not in fields and "outlet" in fields:
            fields["title"] = fields.pop("outlet")
        if "title" not in fields:
            warnings.append(f"Sheet '{sheet_name}' has no recognized title column and was skipped.")
            continue
        rich_cells = rich_by_sheet.get(sheet_name, {})
        for row in range(header_row + 1, sheet.max_row + 1):
            title = clean_text(_value(sheet, row, fields, "title"))
            author_text = clean_text(_value(sheet, row, fields, "authors"))
            if not title and not author_text:
                continue
            selected, bold, method = _select_author(
                publication_type, author_text, _rich_at(rich_cells, row, fields.get("authors"))
            )
            date_value = _value(sheet, row, fields, "date")
            year, publication_date, date_status = parse_year(date_value)
            indexing = clean_text(_value(sheet, row, fields, "indexing"))
            flags = indexing_flags(indexing)
            quartile = normalize_quartile(_value(sheet, row, fields, "quartile")) if publication_type == "Journal" else ""
            doi = clean_text(_value(sheet, row, fields, "doi"))
            outlet = clean_text(_value(sheet, row, fields, "outlet"))
            volume_issue = clean_text(_value(sheet, row, fields, "volume_issue"))
            issue = clean_text(_value(sheet, row, fields, "issue"))
            if issue and issue not in volume_issue:
                volume_issue = " / ".join(filter(None, [volume_issue, issue]))
            record: dict[str, object] = {
                "Publication ID": publication_id(sheet_name, row, title, doi),
                "Faculty Name": selected,
                "Original Faculty Name": selected,
                "Mapping Status": "Pending",
                "Calendar Year": year,
                "Publication Date": publication_date,
                "Date Status": date_status,
                "Publication Type": publication_type,
                "Title": title,
                "Journal / Outlet Name": outlet,
                "DOI": doi,
                "Publisher": clean_text(_value(sheet, row, fields, "publisher")),
                "Volume / Issue": volume_issue,
                "ISBN/ISSN": clean_text(_value(sheet, row, fields, "isbn")),
                "Indexing Raw": indexing,
                "Quartile": quartile,
                "Q1_Flag": int(quartile == "Q1"),
                "Q2_Flag": int(quartile == "Q2"),
                "Q3_Flag": int(quartile == "Q3"),
                "Q4_Flag": int(quartile == "Q4"),
                "Impact Factor": _safe_float(_value(sheet, row, fields, "impact_factor")),
                "Original Authors": author_text,
                "First Bold Author": bold,
                "Author Selection Method": method,
                "Source Sheet": sheet_name,
                "Source Row": row,
                **flags,
            }
            records.append(record)

    expected = {"Journal", "Conference", "Book Chapter", "Book"}
    missing = sorted(expected - detected_types)
    if missing:
        warnings.append("Missing optional publication sheets: " + ", ".join(missing))
    frame = pd.DataFrame.from_records(records)
    if frame.empty:
        frame = pd.DataFrame(columns=MASTER_COLUMNS)
    else:
        for column in MASTER_COLUMNS:
            if column not in frame:
                frame[column] = ""
        frame = frame[MASTER_COLUMNS]
    metadata = {
        "detected_sheets": detected,
        "all_sheets": workbook.sheetnames,
        "warnings": warnings,
        "processed_rows": len(frame),
        "years": sorted(int(year) for year in frame["Calendar Year"].dropna().unique()),
    }
    return frame, metadata
